from __future__ import annotations

import csv
import re
import tempfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE_DIR = Path(__file__).resolve().parent
REPORT_ASSETS_DIR = BASE_DIR / "assets" / "report-kit"
DEFAULT_COVER_BACKGROUND = REPORT_ASSETS_DIR / "cover-background.png"
DEFAULT_SLIDE_BACKGROUND = REPORT_ASSETS_DIR / "slide-background.png"
DEFAULT_REPORT_LOGO = REPORT_ASSETS_DIR / "logo-prenseable.png"
OUTPUT_DIR = BASE_DIR / "output"
FONT_FAMILY = "Open Sans"
TITLE_SIZE = 26
BODY_SIZE = 12

COLOR_PRIMARY = RGBColor(0xFF, 0x40, 0xB4)
COLOR_SECONDARY = RGBColor(0xEC, 0xEC, 0xEC)
COLOR_ACCENT = RGBColor(0xDE, 0x0A, 0x98)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_DARK = RGBColor(0x23, 0x22, 0x27)
COLOR_MUTED = RGBColor(0x6D, 0x6A, 0x73)
COLOR_BORDER = RGBColor(0xD9, 0xD5, 0xDD)
CHART_COLORS = [
    RGBColor(0xFF, 0x40, 0xB4),
    RGBColor(0xDE, 0x0A, 0x98),
    RGBColor(0x8F, 0x6E, 0x84),
    RGBColor(0x6F, 0x8A, 0x93),
    RGBColor(0xE5, 0xB8, 0xD4),
    RGBColor(0xC7, 0xD8, 0xDE),
]


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".txt":
            return path.read_text(encoding="utf-8")
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
    except UnicodeDecodeError as exc:
        raise ValueError("El archivo TXT no esta en UTF-8 o no se pudo leer correctamente.") from exc
    except Exception as exc:
        raise ValueError(f"No se pudo leer el archivo {path.name}. Revisa que sea un {suffix} valido.") from exc
    raise ValueError(f"Formato no soportado: {suffix}")


def get_output_dir() -> Path:
    candidates = [
        OUTPUT_DIR,
        Path(tempfile.gettempdir()) / "strategic-ppt-generator-output",
    ]
    for candidate in candidates:
        try:
            candidate.mkdir(parents=True, exist_ok=True)
            probe = candidate / ".write_test"
            probe.write_text("ok", encoding="utf-8")
            probe.unlink(missing_ok=True)
            return candidate
        except OSError:
            continue
    fallback = Path(tempfile.mkdtemp(prefix="strategic-ppt-generator-"))
    fallback.mkdir(parents=True, exist_ok=True)
    return fallback

MONTH_ALIASES = {
    "jan": "Ene",
    "january": "Ene",
    "ene": "Ene",
    "enero": "Ene",
    "feb": "Feb",
    "february": "Feb",
    "febrero": "Feb",
    "mar": "Mar",
    "march": "Mar",
    "marzo": "Mar",
    "apr": "Abr",
    "april": "Abr",
    "abr": "Abr",
    "abril": "Abr",
    "may": "May",
    "mayo": "May",
    "jun": "Jun",
    "june": "Jun",
    "junio": "Jun",
    "jul": "Jul",
    "july": "Jul",
    "julio": "Jul",
    "aug": "Ago",
    "august": "Ago",
    "ago": "Ago",
    "agosto": "Ago",
    "sep": "Sep",
    "sept": "Sep",
    "september": "Sep",
    "septiembre": "Sep",
    "oct": "Oct",
    "october": "Oct",
    "octubre": "Oct",
    "nov": "Nov",
    "november": "Nov",
    "noviembre": "Nov",
    "dec": "Dic",
    "december": "Dic",
    "dic": "Dic",
    "diciembre": "Dic",
}


@dataclass
class MetricPoint:
    label: str
    value: float
    unit: str = ""
    raw_value: str = ""
    context: str = ""


@dataclass
class ChartRequest:
    title: str
    chart_type: str
    aliases: tuple[str, ...]
    metrics: list[MetricPoint] = field(default_factory=list)
    page_number: int | None = None
    matched_alias: str = ""


@dataclass
class ReportRow:
    date_text: str
    medium: str
    media_type: str
    tier: str
    valuation: str = ""
    reach: str = ""
    link: str = ""


@dataclass
class ReportData:
    source_text: str
    title: str
    source_name: str
    client_name: str = ""
    report_month: str = ""
    executive_comment: str = ""
    next_steps: str = ""
    rows: list[ReportRow] = field(default_factory=list)
    metrics: list[MetricPoint] = field(default_factory=list)
    page_metrics: list[list[MetricPoint]] = field(default_factory=list)
    requested_charts: list[ChartRequest] = field(default_factory=list)
    monthly_trend: list[MetricPoint] = field(default_factory=list)
    bullets: list[str] = field(default_factory=list)
    evidence: list[str] = field(default_factory=list)
    extraction_notes: list[str] = field(default_factory=list)


@dataclass
class ManualReportInput:
    title: str
    client_name: str
    report_month: str
    monthly_summary: str
    executive_comment: str
    next_steps: str
    tier_metrics: list[MetricPoint]
    media_metrics: list[MetricPoint]
    monthly_trend: list[MetricPoint]
    reach_value: str
    valuation_value: str
    table_rows: list[ReportRow] = field(default_factory=list)


def clean_line(line: str) -> str:
    return re.sub(r"\s+", " ", line.replace("\uf0b7", "-")).strip()


def normalize_text(value: str) -> str:
    value = value.lower()
    replacements = str.maketrans(
        {
            "á": "a",
            "é": "e",
            "í": "i",
            "ó": "o",
            "ú": "u",
            "ñ": "n",
        }
    )
    value = value.translate(replacements)
    return re.sub(r"[^a-z0-9% ]+", " ", value)


def safe_metric_label(label: str) -> str:
    label = re.sub(r"^[\-\*\u2022]+\s*", "", label).strip(" :-")
    label = re.sub(r"\s+", " ", label)
    return label[:40] or "Metrica"


def parse_numeric_token(token: str) -> float | None:
    current = token.strip().replace(" ", "")
    if not re.search(r"\d", current):
        return None
    if "," in current and "." in current:
        current = current.replace(".", "").replace(",", ".")
    elif "," in current:
        current = current.replace(".", "").replace(",", ".")
    try:
        return float(current)
    except ValueError:
        return None


def detect_unit(raw: str) -> str:
    if "%" in raw:
        return "%"
    if re.search(r"\b(?:usd|clp|uf|\$)\b", raw.lower()) or "$" in raw:
        return "$"
    return ""


def metric_sort_key(metric: MetricPoint) -> tuple[int, float]:
    weight = 1 if metric.unit == "%" else 2 if metric.unit == "$" else 0
    return (weight, abs(metric.value))


def normalize_month(label: str) -> str | None:
    key = re.sub(r"[^a-zA-Z]", "", label).lower()
    return MONTH_ALIASES.get(key)


def extract_candidate_metrics(text: str) -> list[MetricPoint]:
    metrics: list[MetricPoint] = []
    seen: set[tuple[str, str]] = set()
    lines = [clean_line(line) for line in text.splitlines()]
    numeric_pattern = re.compile(r"([-+]?\d[\d\.,]*\s*%?)")

    for line in lines:
        if len(line) < 3 or len(line) > 130:
            continue
        matches = list(numeric_pattern.finditer(line))
        if not matches:
            continue
        if len(matches) > 3:
            continue

        first = matches[-1]
        raw_value = first.group(1).strip()
        value = parse_numeric_token(raw_value.replace("%", ""))
        if value is None:
            continue

        prefix = line[: first.start()].strip(" :-|")
        suffix = line[first.end() :].strip(" :-|")
        label = safe_metric_label(prefix or suffix or "Metrica")
        if len(label) < 3:
            continue
        if re.fullmatch(r"\d{1,4}", label):
            continue

        candidate = MetricPoint(
            label=label,
            value=value,
            unit=detect_unit(raw_value),
            raw_value=raw_value,
            context=line,
        )
        fingerprint = (candidate.label.lower(), candidate.raw_value)
        if fingerprint in seen:
            continue
        seen.add(fingerprint)
        metrics.append(candidate)

    return sorted(metrics, key=metric_sort_key, reverse=True)


def build_quant_bullets(metrics: list[MetricPoint]) -> list[str]:
    bullets: list[str] = []
    if not metrics:
        return [
            "No se encontraron metricas numericas suficientemente limpias en el PDF para construir lectura automatica.",
            "El flujo actual funciona mejor cuando el PDF trae texto seleccionable, tablas o etiquetas numericas visibles.",
        ]

    top = metrics[0]
    bullets.append(f"La metrica mas fuerte detectada es {top.label} con {top.raw_value}.")

    percent_metrics = [metric for metric in metrics if metric.unit == "%"]
    if percent_metrics:
        best_percent = max(percent_metrics, key=lambda metric: metric.value)
        bullets.append(f"Entre los porcentajes extraidos, destaca {best_percent.label} con {best_percent.raw_value}.")

    comparable = [metric for metric in metrics if metric.unit == top.unit][:4]
    if len(comparable) >= 2:
        highest = max(comparable, key=lambda metric: metric.value)
        lowest = min(comparable, key=lambda metric: metric.value)
        if highest.label != lowest.label:
            gap = highest.value - lowest.value
            gap_text = f"{gap:.1f}".rstrip("0").rstrip(".")
            suffix = highest.unit if highest.unit else ""
            bullets.append(
                f"La brecha entre la mejor y la menor lectura comparable es {gap_text}{suffix}, entre {highest.label} y {lowest.label}."
            )

    if len(metrics) >= 4:
        bullets.append(
            f"El PDF entrego al menos {len(metrics)} puntos numericos reutilizables, suficiente para una lectura cuantitativa inicial."
        )
    else:
        bullets.append("La lectura cuantitativa es preliminar; conviene complementar con mas datos estructurados si quieren mayor precision.")

    return bullets[:4]


def build_evidence(text: str, metrics: list[MetricPoint]) -> list[str]:
    evidence = [metric.context for metric in metrics[:6] if metric.context]
    if evidence:
        return evidence
    lines = [clean_line(line) for line in text.splitlines() if clean_line(line)]
    return lines[:6]


def extract_pdf_pages(pdf_path: Path) -> list[str]:
    reader = PdfReader(str(pdf_path))
    return [(page.extract_text() or "").strip() for page in reader.pages]


def default_chart_requests() -> list[ChartRequest]:
    return [
        ChartRequest(
            title="Proporcion de Datos",
            chart_type="pie",
            aliases=("proporcion de datos", "proporcion de tiers", "tiers"),
        ),
        ChartRequest(
            title="Proporcion Tipo de Contenido",
            chart_type="pie",
            aliases=("proporcion tipo de contenido", "proporcion tipo de contenid", "proporcion tipo comunicado"),
        ),
        ChartRequest(
            title="Tipos de Medios",
            chart_type="pie",
            aliases=("tipos de medios", "tipo de medios"),
        ),
        ChartRequest(
            title="Medios",
            chart_type="bar",
            aliases=("medios",),
        ),
    ]


def title_case_words(value: str) -> str:
    return " ".join(part.capitalize() for part in clean_line(value).split())


def extract_report_rows(text: str) -> list[ReportRow]:
    rows: list[ReportRow] = []
    pattern = re.compile(
        r"^(?P<date>\d{1,2}\s+\w+\s+\S+)\s+"
        r"(?P<client>\S+)\s+"
        r"(?P<medium>.+?)\s+"
        r"(?P<media_type>Digital|Impreso|Radio|TV|Television|Televisión)\s+"
        r"(?P<tier>Tier\s+\d+)\s+"
        r"(?P<communication_type>.+?)\s+https?://",
        re.IGNORECASE,
    )

    for raw_line in text.splitlines():
        line = clean_line(raw_line)
        match = pattern.search(line)
        if not match:
            continue
        rows.append(
            ReportRow(
                date_text=match.group("date"),
                client=match.group("client"),
                medium=clean_line(match.group("medium")),
                media_type=title_case_words(match.group("media_type")),
                tier=title_case_words(match.group("tier")),
                communication_type=title_case_words(match.group("communication_type")),
            )
        )
    return rows


def build_metrics_from_counter(items: list[str]) -> list[MetricPoint]:
    counter: dict[str, int] = {}
    for item in items:
        key = clean_line(item)
        if not key:
            continue
        counter[key] = counter.get(key, 0) + 1
    ordered = sorted(counter.items(), key=lambda item: (-item[1], item[0].lower()))
    return [MetricPoint(label=label, value=int(value), raw_value=str(value)) for label, value in ordered]


def chart_requests_from_rows(rows: list[ReportRow]) -> list[ChartRequest]:
    requests = default_chart_requests()
    if not rows:
        return requests

    mappings = {
        "Proporcion de Datos": build_metrics_from_counter([row.tier for row in rows]),
        "Proporcion Tipo de Contenido": build_metrics_from_counter([row.communication_type for row in rows]),
        "Tipos de Medios": build_metrics_from_counter([row.media_type for row in rows]),
        "Medios": build_metrics_from_counter([row.medium for row in rows]),
    }
    resolved: list[ChartRequest] = []
    for request in requests:
        resolved.append(
            ChartRequest(
                title=request.title,
                chart_type=request.chart_type,
                aliases=request.aliases,
                metrics=mappings.get(request.title, []),
                page_number=1 if mappings.get(request.title) else None,
                matched_alias=request.aliases[0],
            )
        )
    return resolved


def parse_table_rows(text: str) -> list[ReportRow]:
    rows: list[ReportRow] = []
    for raw_line in text.splitlines():
        line = clean_line(raw_line)
        if not line:
            continue
        parts = [part.strip() for part in line.split("|")]
        if len(parts) < 6:
            continue
        rows.append(
            ReportRow(
                date_text=parts[0],
                medium=parts[1],
                tier=parts[2] if len(parts) > 2 else "",
                media_type=parts[3] if len(parts) > 3 else "",
                valuation=parts[4] if len(parts) > 4 else "",
                reach=parts[5] if len(parts) > 5 else "",
                link=parts[6] if len(parts) > 6 else "",
            )
        )
    return rows


def parse_table_file(path: Path) -> list[ReportRow]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            return rows_from_dicts(list(reader))
    if suffix == ".xlsx":
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(value).strip() if value is not None else "" for value in values[0]]
        entries: list[dict[str, str]] = []
        for row in values[1:]:
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            entry = {headers[index]: ("" if index >= len(row) or row[index] is None else str(row[index]).strip()) for index in range(len(headers))}
            entries.append(entry)
        return rows_from_dicts(entries)
    return []


def rows_from_dicts(entries: list[dict[str, str]]) -> list[ReportRow]:
    rows: list[ReportRow] = []
    for entry in entries:
        normalized = {normalize_text(str(key)).strip(): str(value).strip() for key, value in entry.items()}
        date_text = normalized.get("fecha") or normalized.get("date") or normalized.get("mes") or ""
        medium = normalized.get("medio") or normalized.get("media") or ""
        media_type = normalized.get("tipo medio") or normalized.get("tipo de medio") or normalized.get("tipo_medio") or ""
        tier = normalized.get("tier") or ""
        valuation = normalized.get("valorizacion") or normalized.get("valorización") or normalized.get("valor estimado") or normalized.get("valor_estimado") or ""
        reach = normalized.get("alcance") or normalized.get("alcance estimado") or normalized.get("alcance_estimado") or ""
        link = normalized.get("link") or normalized.get("enlace") or normalized.get("url") or ""
        if not any([date_text, medium, media_type, tier, valuation, reach, link]):
            continue
        rows.append(
            ReportRow(
                date_text=date_text,
                medium=medium,
                media_type=media_type,
                tier=tier,
                valuation=valuation,
                reach=reach,
                link=link,
            )
        )
    return rows


def parse_section_metrics(lines: list[str], alias: str) -> list[MetricPoint]:
    metrics: list[MetricPoint] = []
    seen: set[tuple[str, str]] = set()
    numeric_pattern = re.compile(r"([-+]?\d[\d\.,]*\s*%?)")
    alias_norm = normalize_text(alias).strip()

    for index, raw_line in enumerate(lines):
        line = clean_line(raw_line)
        if not line:
            continue
        numbers = list(numeric_pattern.finditer(line))
        if not numbers:
            continue

        last = numbers[-1]
        raw_value = last.group(1).strip()
        value = parse_numeric_token(raw_value.replace("%", ""))
        if value is None:
            continue

        prefix = clean_line(line[: last.start()].strip(" :-|"))
        suffix = clean_line(line[last.end() :].strip(" :-|"))
        label = prefix or suffix

        if not label or normalize_text(label).strip() == alias_norm:
            if index > 0:
                previous = clean_line(lines[index - 1])
                if previous and not numeric_pattern.search(previous):
                    label = previous
        if not label:
            label = f"Categoria {len(metrics) + 1}"

        label = safe_metric_label(label)
        if len(label) < 2:
            continue

        fingerprint = (label.lower(), raw_value)
        if fingerprint in seen:
            continue
        seen.add(fingerprint)
        metrics.append(
            MetricPoint(
                label=label,
                value=value,
                unit=detect_unit(raw_value),
                raw_value=raw_value,
                context=line,
            )
        )

    return metrics[:8]


def extract_target_chart(page_texts: list[str], request: ChartRequest) -> ChartRequest:
    result = ChartRequest(title=request.title, chart_type=request.chart_type, aliases=request.aliases)
    for page_number, page_text in enumerate(page_texts, start=1):
        lines = [line for line in page_text.splitlines() if clean_line(line)]
        normalized_lines = [normalize_text(clean_line(line)) for line in lines]
        for line_index, normalized_line in enumerate(normalized_lines):
            if any(alias in normalized_line for alias in request.aliases):
                window = lines[line_index : line_index + 14]
                metrics = parse_section_metrics(window, request.aliases[0])
                if metrics:
                    result.metrics = metrics
                result.page_number = page_number
                result.matched_alias = request.aliases[0]
                return result
    return result


def extract_monthly_trend(text: str) -> list[MetricPoint]:
    metrics: list[MetricPoint] = []
    lines = [clean_line(line) for line in text.splitlines() if clean_line(line)]
    for line in lines:
        normalized = normalize_text(line)
        for token in normalized.split():
            month = normalize_month(token)
            if not month:
                continue
            numbers = re.findall(r"([-+]?\d[\d\.,]*)", line)
            if not numbers:
                continue
            value = parse_numeric_token(numbers[-1])
            if value is None:
                continue
            metrics.append(
                MetricPoint(
                    label=month,
                    value=value,
                    raw_value=numbers[-1],
                    context=line,
                )
            )
            break

    if metrics:
        deduped: dict[str, MetricPoint] = {}
        for metric in metrics:
            deduped[metric.label] = metric
        ordered_months = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
        result = [deduped[month] for month in ordered_months if month in deduped]
        if len(result) == 1:
            placeholders = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
            only = result[0]
            return [
                MetricPoint(label=month, value=int(only.value) if month == only.label else 0, raw_value=str(int(only.value)) if month == only.label else "0")
                for month in placeholders
            ]
        return result

    placeholders = [("Ene", 1), ("Feb", 2), ("Mar", 3), ("Abr", 4), ("May", 5), ("Jun", 6)]
    return [MetricPoint(label=label, value=int(value), raw_value=str(value)) for label, value in placeholders]


def extract_monthly_trend_from_rows(rows: list[ReportRow]) -> list[MetricPoint]:
    if not rows:
        return extract_monthly_trend("")

    counter: dict[str, int] = {}
    for row in rows:
        parts = row.date_text.split()
        if len(parts) < 2:
            continue
        month = normalize_month(parts[1]) or title_case_words(parts[1])[:3]
        counter[month] = counter.get(month, 0) + 1
    ordered_months = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
    metrics = [MetricPoint(label=month, value=int(counter[month]), raw_value=str(counter[month])) for month in ordered_months if month in counter]
    return metrics or extract_monthly_trend("")


def parse_report_pdf(
    pdf_path: Path,
    executive_comment: str = "",
    next_steps: str = "",
    report_title: str | None = None,
    client_name: str = "",
    report_month: str = "",
) -> ReportData:
    text = extract_text(pdf_path)
    page_texts = extract_pdf_pages(pdf_path)
    rows = extract_report_rows(text)
    metrics = extract_candidate_metrics(text)
    page_metrics = [extract_candidate_metrics(page_text) for page_text in page_texts]
    requested_charts = chart_requests_from_rows(rows)
    if not any(chart.metrics for chart in requested_charts):
        requested_charts = [extract_target_chart(page_texts, request) for request in default_chart_requests()]
    monthly_trend = extract_monthly_trend_from_rows(rows)
    useful_pages = sum(1 for page in page_metrics if page)
    matched_charts = sum(1 for chart in requested_charts if chart.metrics)
    extraction_notes = []
    if metrics:
        extraction_notes.append(
            f"Se detectaron {len(metrics)} metricas candidatas desde texto legible del PDF."
        )
    else:
        extraction_notes.append(
            "No se detectaron metricas legibles; si el grafico esta embebido como imagen, sera necesario OCR o carga manual."
        )
    extraction_notes.append(
        f"Se detectaron {len(page_texts)} paginas y {useful_pages} con datos numericos reutilizables para slides individuales."
    )
    extraction_notes.append(
        f"Se identificaron {matched_charts} de 4 graficos objetivo para construir slides dedicadas."
    )
    if rows:
        extraction_notes.append(f"Se reconstruyeron {len(rows)} publicaciones desde la tabla principal del PDF.")

    title = report_title or f"Reporte Automatico {datetime.now():%B %Y}"
    return ReportData(
        source_text=text,
        title=title,
        source_name=pdf_path.name,
        client_name=client_name.strip(),
        report_month=report_month.strip(),
        executive_comment=executive_comment.strip(),
        next_steps=next_steps.strip(),
        rows=rows,
        metrics=metrics,
        page_metrics=page_metrics,
        requested_charts=requested_charts,
        monthly_trend=monthly_trend,
        bullets=build_quant_bullets(metrics),
        evidence=build_evidence(text, metrics),
        extraction_notes=extraction_notes,
    )


def add_picture_safe(slide, path: Path, left, top, width=None, height=None):
    if not path or not path.exists():
        return None
    kwargs = {}
    if width is not None:
        kwargs["width"] = width
    if height is not None:
        kwargs["height"] = height
    return slide.shapes.add_picture(str(path), left, top, **kwargs)


def resolve_asset_path(candidate: Path | None, default: Path | None) -> Path | None:
    if candidate and candidate.exists():
        return candidate
    if default and default.exists():
        return default
    return None


def apply_report_background(slide, background_path: Path | None) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_SECONDARY

    resolved_background = resolve_asset_path(background_path, DEFAULT_SLIDE_BACKGROUND)
    if resolved_background:
        add_picture_safe(slide, resolved_background, 0, 0, width=Inches(13.333), height=Inches(7.5))
    else:
        top_band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.0))
        top_band.fill.solid()
        top_band.fill.fore_color.rgb = COLOR_PRIMARY
        top_band.line.fill.background()
        accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.35), Inches(0.35), Inches(12.6), Inches(0.03))
        accent.fill.solid()
        accent.fill.fore_color.rgb = COLOR_ACCENT
        accent.line.fill.background()


def add_logo_box(slide, logo_path: Path | None) -> None:
    resolved_logo = resolve_asset_path(logo_path, DEFAULT_REPORT_LOGO)
    if resolved_logo:
        add_picture_safe(slide, resolved_logo, Inches(11.5), Inches(0.18), width=Inches(1.35))
    else:
        box = slide.shapes.add_textbox(Inches(11.5), Inches(0.22), Inches(1.35), Inches(0.45))
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = "LOGO"
        run.font.name = FONT_FAMILY
        run.font.size = Pt(BODY_SIZE)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY


def add_slide_base(prs: Presentation, title: str, background_path: Path | None, logo_path: Path | None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_report_background(slide, background_path)
    add_logo_box(slide, logo_path)

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.28), Inches(9.8), Inches(0.6))
    p = title_box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title
    run.font.name = FONT_FAMILY
    run.font.bold = True
    run.font.size = Pt(TITLE_SIZE)
    run.font.color.rgb = COLOR_PRIMARY
    return slide


def add_text_block(slide, title: str, body: str, left, top, width, height, accent_color: RGBColor = COLOR_PRIMARY) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_WHITE
    shape.line.color.rgb = COLOR_BORDER
    tf = shape.text_frame
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.14)
    tf.margin_top = Inches(0.1)
    tf.margin_bottom = Inches(0.08)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_FAMILY
    p.font.bold = True
    p.font.size = Pt(TITLE_SIZE)
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(8)
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(BODY_SIZE)
    p2.font.color.rgb = COLOR_MUTED
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, Inches(0.14), height)
    accent.fill.solid()
    accent.fill.fore_color.rgb = accent_color
    accent.line.fill.background()


def add_bullet_list(slide, title: str, bullets: list[str], left, top, width, height) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_WHITE
    shape.line.color.rgb = COLOR_BORDER
    tf = shape.text_frame
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.14)
    tf.margin_top = Inches(0.1)
    tf.margin_bottom = Inches(0.08)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_FAMILY
    p.font.bold = True
    p.font.size = Pt(TITLE_SIZE)
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(8)
    for bullet in bullets:
        item = tf.add_paragraph()
        item.text = f"- {bullet}"
        item.level = 0
        item.font.name = FONT_FAMILY
        item.font.size = Pt(BODY_SIZE)
        item.font.color.rgb = COLOR_MUTED
        item.space_before = Pt(2)


def add_metric_cards(slide, metrics: list[MetricPoint]) -> None:
    cards = metrics[:4]
    while len(cards) < 4:
        cards.append(MetricPoint(label="Dato pendiente", value=0, raw_value="[sin dato]"))

    for index, metric in enumerate(cards):
        left = Inches(0.85 + (index % 2) * 6.0)
        top = Inches(1.45 + (index // 2) * 1.7)
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, Inches(5.5), Inches(1.35))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_WHITE
        shape.line.color.rgb = COLOR_BORDER
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = metric.label
        p.font.name = FONT_FAMILY
        p.font.size = Pt(BODY_SIZE)
        p.font.color.rgb = COLOR_MUTED
        p2 = tf.add_paragraph()
        p2.text = metric.raw_value or "[sin dato]"
        p2.font.name = FONT_FAMILY
        p2.font.bold = True
        p2.font.size = Pt(20)
        p2.font.color.rgb = COLOR_PRIMARY if index % 2 == 0 else COLOR_ACCENT


def add_metric_chart(slide, metrics: list[MetricPoint]) -> None:
    chart_metrics = [metric for metric in metrics if metric.value >= 0][:6]
    if not chart_metrics:
        add_text_block(
            slide,
            "Datos cuantitativos",
            "No hubo suficientes valores consistentes para graficar automaticamente desde el PDF.",
            Inches(0.85),
            Inches(1.55),
            Inches(11.7),
            Inches(1.2),
        )
        return

    chart_data = CategoryChartData()
    chart_data.categories = [metric.label[:18] for metric in chart_metrics]
    chart_data.add_series("Valor", [metric.value for metric in chart_metrics])

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.9),
        Inches(1.7),
        Inches(7.1),
        Inches(4.5),
        chart_data,
    ).chart
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = True
    chart.category_axis.tick_labels.font.size = Pt(BODY_SIZE)
    chart.value_axis.tick_labels.font.size = Pt(BODY_SIZE)
    series = chart.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = COLOR_PRIMARY
    series.format.line.color.rgb = COLOR_ACCENT
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END

    add_bullet_list(
        slide,
        "Lectura cuantitativa",
        build_quant_bullets(metrics),
        Inches(8.25),
        Inches(1.7),
        Inches(4.2),
        Inches(4.5),
    )


def add_distribution_chart(
    slide,
    metrics: list[MetricPoint],
    chart_type: str,
    left,
    top,
    width,
    height,
):
    if not metrics:
        return
    chart_data = CategoryChartData()
    chart_data.categories = [metric.label[:24] for metric in metrics]
    chart_data.add_series("Valor", [int(metric.value) for metric in metrics])
    ppt_chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED
    chart = slide.shapes.add_chart(ppt_chart_type, left, top, width, height, chart_data).chart
    chart.has_title = False
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = True
    chart.category_axis.tick_labels.font.size = Pt(BODY_SIZE)
    chart.value_axis.tick_labels.font.size = Pt(BODY_SIZE)
    chart.category_axis.tick_labels.offset = 100
    series = chart.series[0]
    chart.plots[0].vary_by_categories = True
    for index, point in enumerate(series.points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = CHART_COLORS[index % len(CHART_COLORS)]
        point.format.line.color.rgb = COLOR_WHITE
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.show_value = True


def build_chart_specific_comments(chart: ChartRequest) -> list[str]:
    if not chart.metrics:
        return [
            "No se pudo reconstruir este grafico con suficiente precision desde el PDF.",
            "Si el bloque viene como imagen, conviene complementar con OCR o carga manual de datos.",
        ]
    top = max(chart.metrics, key=lambda metric: metric.value)
    total = sum(metric.value for metric in chart.metrics) or 1
    use_percentages = chart.title in {"Distribución de Tiers", "Distribución de Medios"}
    if use_percentages:
        top_pct = round((top.value / total) * 100)
        comments = [f"La categoria dominante es {top.label} con {top_pct}%."]
    else:
        comments = [f"La categoria dominante es {top.label} con {top.raw_value}."]
    if len(chart.metrics) >= 2:
        ordered = sorted(chart.metrics, key=lambda metric: metric.value, reverse=True)
        second = ordered[1]
        if use_percentages:
            second_pct = round((second.value / total) * 100)
            gap_pct = round(((ordered[0].value - ordered[1].value) / total) * 100)
            comments.append(f"La segunda lectura mas relevante es {second.label} con {second_pct}%.")
            comments.append(f"La brecha entre ambas primeras categorias es {gap_pct} puntos porcentuales.")
        else:
            comments.append(f"La segunda lectura mas relevante es {second.label} con {second.raw_value}.")
            gap = ordered[0].value - ordered[1].value
            gap_text = f"{gap:.1f}".rstrip("0").rstrip(".")
            suffix = ordered[0].unit if ordered[0].unit else ""
            comments.append(f"La brecha entre ambas primeras categorias es {gap_text}{suffix}.")
    return comments[:4]


def add_named_chart_slide(
    prs: Presentation,
    chart: ChartRequest,
    background_path: Path | None,
    logo_path: Path | None,
) -> None:
    slide = add_slide_base(prs, chart.title, background_path, logo_path)
    if not chart.metrics:
        add_text_block(
            slide,
            "Grafico pendiente",
            "No se pudo leer este grafico con claridad desde el PDF. La slide queda lista para que luego se reemplacen los datos manualmente si hace falta.",
            Inches(0.85),
            Inches(1.55),
            Inches(11.55),
            Inches(1.45),
        )
        add_bullet_list(
            slide,
            "Comentario cuantitativo",
            build_chart_specific_comments(chart),
            Inches(0.85),
            Inches(3.25),
            Inches(11.55),
            Inches(2.2),
        )
        return

    add_distribution_chart(
        slide,
        chart.metrics[:6],
        chart.chart_type,
        Inches(0.85),
        Inches(1.55),
        Inches(6.6),
        Inches(4.5),
    )
    add_bullet_list(
        slide,
        "Comentario cuantitativo",
        build_chart_specific_comments(chart),
        Inches(7.8),
        Inches(1.55),
        Inches(4.65),
        Inches(2.35),
    )
    add_bullet_list(
        slide,
        "Datos rescatados",
        [f"{metric.label}: {metric.raw_value}" for metric in chart.metrics[:6]],
        Inches(7.8),
        Inches(4.1),
        Inches(4.65),
        Inches(1.95),
    )


def add_monthly_trend_slide(
    prs: Presentation,
    trend_metrics: list[MetricPoint],
    background_path: Path | None,
    logo_path: Path | None,
) -> None:
    slide = add_slide_base(prs, "Publicaciones Mes a Mes", background_path, logo_path)
    metrics = trend_metrics[:12]
    if len(metrics) < 12:
        ordered = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
        current = {metric.label: metric for metric in metrics}
        metrics = [
            current.get(month, MetricPoint(label=month, value=0, raw_value="0"))
            for month in ordered
        ]
    chart_data = CategoryChartData()
    chart_data.categories = [metric.label for metric in metrics]
    chart_data.add_series("Publicaciones", [int(metric.value) for metric in metrics])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE,
        Inches(0.85),
        Inches(1.55),
        Inches(7.1),
        Inches(4.45),
        chart_data,
    ).chart
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = True
    chart.category_axis.tick_labels.font.size = Pt(BODY_SIZE)
    chart.value_axis.tick_labels.font.size = Pt(BODY_SIZE)
    series = chart.series[0]
    series.format.line.color.rgb = COLOR_PRIMARY

    add_text_block(
        slide,
        "Lectura de avance",
        "",
        Inches(8.15),
        Inches(1.55),
        Inches(4.3),
        Inches(2.4),
    )


def add_page_analysis_slide(
    prs: Presentation,
    page_number: int,
    metrics: list[MetricPoint],
    background_path: Path | None,
    logo_path: Path | None,
) -> None:
    slide = add_slide_base(prs, f"Grafico {page_number}", background_path, logo_path)
    if not metrics:
        add_text_block(
            slide,
            "Lectura de la pagina",
            "No se rescataron metricas numericas legibles en esta pagina. Si el grafico viene como imagen, conviene complementar con OCR o carga manual.",
            Inches(0.85),
            Inches(1.55),
            Inches(11.55),
            Inches(1.4),
        )
        return

    chart_metrics = metrics[:6]
    chart_data = CategoryChartData()
    chart_data.categories = [metric.label[:18] for metric in chart_metrics]
    chart_data.add_series("Valor", [metric.value for metric in chart_metrics])

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.9),
        Inches(1.7),
        Inches(6.9),
        Inches(4.3),
        chart_data,
    ).chart
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = True
    chart.category_axis.tick_labels.font.size = Pt(BODY_SIZE)
    chart.value_axis.tick_labels.font.size = Pt(BODY_SIZE)
    series = chart.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = COLOR_PRIMARY
    series.format.line.color.rgb = COLOR_ACCENT
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END

    add_bullet_list(
        slide,
        "Opinion cuantitativa",
        build_quant_bullets(metrics),
        Inches(8.0),
        Inches(1.7),
        Inches(4.45),
        Inches(2.7),
    )
    add_bullet_list(
        slide,
        "Datos rescatados",
        [f"{metric.label}: {metric.raw_value}" for metric in metrics[:5]],
        Inches(8.0),
        Inches(4.6),
        Inches(4.45),
        Inches(1.55),
    )


def add_cover_slide(prs: Presentation, data: ReportData, background_path: Path | None, logo_path: Path | None) -> None:
    cover_background = DEFAULT_COVER_BACKGROUND if DEFAULT_COVER_BACKGROUND.exists() else resolve_asset_path(None, DEFAULT_COVER_BACKGROUND)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_PRIMARY
    if cover_background:
        add_picture_safe(slide, cover_background, 0, 0, width=Inches(13.333), height=Inches(7.5))

    title_box = slide.shapes.add_textbox(Inches(2.0), Inches(2.6), Inches(9.33), Inches(0.5))
    title_p = title_box.text_frame.paragraphs[0]
    title_p.alignment = PP_ALIGN.CENTER
    title_run = title_p.add_run()
    title_run.text = data.title or "Reporte Automatico"
    title_run.font.name = FONT_FAMILY
    title_run.font.bold = True
    title_run.font.size = Pt(36)
    title_run.font.color.rgb = COLOR_WHITE

    client_box = slide.shapes.add_textbox(Inches(2.0), Inches(3.2), Inches(9.33), Inches(0.5))
    p = client_box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = data.client_name or "[Nombre del cliente]"
    run.font.name = FONT_FAMILY
    run.font.bold = True
    run.font.size = Pt(36)
    run.font.color.rgb = COLOR_WHITE

    month_box = slide.shapes.add_textbox(Inches(2.0), Inches(3.8), Inches(9.33), Inches(0.45))
    p2 = month_box.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = data.report_month or datetime.now().strftime("%B %Y")
    run2.font.name = FONT_FAMILY
    run2.font.bold = False
    run2.font.size = Pt(36)
    run2.font.color.rgb = COLOR_WHITE


def add_extraction_slide(prs: Presentation, data: ReportData, background_path: Path | None, logo_path: Path | None) -> None:
    slide = add_slide_base(prs, "Lectura del PDF", background_path, logo_path)
    add_text_block(
        slide,
        "Estado de extraccion",
        "\n".join(data.extraction_notes),
        Inches(0.85),
        Inches(1.45),
        Inches(5.45),
        Inches(1.25),
    )
    add_text_block(
        slide,
        "Capacidad actual",
        "El sistema reutiliza texto, tablas y etiquetas numericas presentes en el PDF. Si un grafico viene solo como imagen, la lectura automatica puede quedar parcial.",
        Inches(6.45),
        Inches(1.45),
        Inches(5.95),
        Inches(1.25),
        accent_color=COLOR_ACCENT,
    )
    add_bullet_list(
        slide,
        "Evidencia detectada",
        data.evidence[:5] or ["Sin evidencia legible para mostrar."],
        Inches(0.85),
        Inches(3.0),
        Inches(11.55),
        Inches(3.2),
    )


def add_exec_slide(prs: Presentation, data: ReportData, background_path: Path | None, logo_path: Path | None) -> None:
    slide = add_slide_base(prs, "Comentario Ejecutivo", background_path, logo_path)
    add_text_block(
        slide,
        "Resumen de gestion",
        data.executive_comment or "[Completar comentario ejecutivo de la gestion del mes]",
        Inches(0.85),
        Inches(1.5),
        Inches(11.55),
        Inches(4.8),
    )


def add_next_steps_slide(prs: Presentation, data: ReportData, background_path: Path | None, logo_path: Path | None) -> None:
    slide = add_slide_base(prs, "Pasos a Seguir", background_path, logo_path)
    bullets = [clean_line(line) for line in data.next_steps.splitlines() if clean_line(line)]
    if not bullets and data.next_steps.strip():
        bullets = [item.strip() for item in re.split(r"[;\n]", data.next_steps) if item.strip()]
    add_bullet_list(
        slide,
        "Mes siguiente",
        bullets or ["[Completar siguientes pasos para el proximo mes]"],
        Inches(0.85),
        Inches(1.45),
        Inches(11.55),
        Inches(4.9),
    )


def add_month_summary_slide(
    prs: Presentation,
    summary: str,
    background_path: Path | None,
    logo_path: Path | None,
) -> None:
    slide = add_slide_base(prs, "Breve Resumen del Mes", background_path, logo_path)
    add_text_block(
        slide,
        "Resumen",
        summary or "[Completar breve resumen del mes]",
        Inches(0.85),
        Inches(1.5),
        Inches(11.55),
        Inches(4.7),
    )


def add_scope_slide(
    prs: Presentation,
    reach_value: str,
    valuation_value: str,
    background_path: Path | None,
    logo_path: Path | None,
) -> None:
    slide = add_slide_base(prs, "Alcance y Valorización", background_path, logo_path)
    cards = [
        ("Alcance de la gestión", reach_value or "[Completar alcance]"),
        ("Valorización", valuation_value or "[Completar valorización]"),
    ]
    for index, (title, value) in enumerate(cards):
        top = Inches(1.95 + index * 2.35)
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.15), top, Inches(4.75), Inches(1.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_WHITE
        shape.line.color.rgb = COLOR_BORDER
        tf = shape.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = title
        p.font.name = FONT_FAMILY
        p.font.bold = True
        p.font.size = Pt(18)
        p.font.color.rgb = COLOR_PRIMARY
        p2 = tf.add_paragraph()
        p2.text = value
        p2.font.name = FONT_FAMILY
        p2.font.bold = True
        p2.font.size = Pt(22)
        p2.font.color.rgb = COLOR_DARK
        p2.alignment = PP_ALIGN.CENTER


def add_results_table_slide(
    prs: Presentation,
    rows: list[ReportRow],
    background_path: Path | None,
    logo_path: Path | None,
) -> None:
    slide = add_slide_base(prs, "Detalle de Resultados", background_path, logo_path)
    headers = ["Fecha", "Medio", "Tier", "Tipo Medio", "Valorización", "Alcance", "Link"]
    lefts = [0.35, 1.75, 4.55, 5.75, 7.25, 9.0, 10.75]
    widths = [1.3, 2.65, 1.0, 1.35, 1.6, 1.5, 1.75]
    for left, width, header in zip(lefts, widths, headers):
        cell = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(left), Inches(1.3), Inches(width), Inches(0.48))
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        cell.line.fill.background()
        tf = cell.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = header
        p.font.name = FONT_FAMILY
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_WHITE

    table_rows = rows[:8] if rows else [ReportRow("[Fecha]", "[Medio]", "[Tipo]", "[Tier]", "[Valorización]", "[Alcance]", "[Link]")]
    for row_index, row in enumerate(table_rows):
        top = Inches(1.8 + row_index * 0.56)
        values = [row.date_text, row.medium, row.tier, row.media_type, row.valuation, row.reach, row.link]
        for col_index, (left, width, value) in enumerate(zip(lefts, widths, values)):
            cell = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(left), top, Inches(width), Inches(0.52))
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_WHITE if row_index % 2 == 0 else COLOR_SECONDARY
            cell.line.color.rgb = COLOR_BORDER
            tf = cell.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.text = "Abrir nota" if col_index == 6 and value else value
            p.font.name = FONT_FAMILY
            p.font.size = Pt(10)
            p.font.color.rgb = COLOR_PRIMARY if col_index == 6 and value else COLOR_DARK
            if col_index == 6 and value and p.runs:
                try:
                    p.runs[0].hyperlink.address = value
                except Exception:
                    pass


def write_report_summary(
    output_path: Path,
    data: ReportData,
    background_path: Path | None,
    logo_path: Path | None,
) -> Path:
    summary_path = output_path.with_suffix(".txt")
    lines = [
        f"Archivo generado: {output_path.name}",
        f"Fuente: {data.source_name}",
        f"Fecha: {datetime.now().isoformat(timespec='seconds')}",
        "",
        "Branding aplicado:",
        f"- Fondo personalizado: {'si' if background_path and background_path.exists() else 'no'}",
        f"- Logo personalizado: {'si' if logo_path and logo_path.exists() else 'no'}",
        "",
        "Lectura cuantitativa:",
    ]
    lines.extend(f"- {bullet}" for bullet in data.bullets)
    lines.append("")
    lines.append("Metricas detectadas:")
    if data.metrics:
        lines.extend(f"- {metric.label}: {metric.raw_value}" for metric in data.metrics[:10])
    else:
        lines.append("- No se detectaron metricas legibles.")
    summary_path.write_text("\n".join(lines), encoding="utf-8")
    return summary_path


def write_manual_report_summary(
    output_path: Path,
    manual_input: ManualReportInput,
    background_path: Path | None,
    logo_path: Path | None,
) -> Path:
    summary_path = output_path.with_suffix(".txt")
    lines = [
        f"Archivo generado: {output_path.name}",
        f"Cliente: {manual_input.client_name}",
        f"Mes: {manual_input.report_month}",
        "",
        "Branding aplicado:",
        f"- Fondo personalizado: {'si' if background_path and background_path.exists() else 'no'}",
        f"- Logo personalizado: {'si' if logo_path and logo_path.exists() else 'no'}",
        "",
        "Slides generadas:",
        "- Breve resumen del mes",
        "- Distribución de tiers",
        "- Distribución de medios",
        "- Publicaciones mes a mes",
        "- Alcance y valorización",
        "- Pasos a seguir",
        "- Tabla de resultados",
    ]
    summary_path.write_text("\n".join(lines), encoding="utf-8")
    return summary_path


def generate_report_from_manual_fields(
    report_title: str,
    client_name: str,
    report_month: str,
    monthly_summary: str,
    executive_comment: str,
    next_steps: str,
    tier_values: dict[str, int],
    media_values: dict[str, int],
    monthly_values: list[int],
    reach_value: str,
    valuation_value: str,
    table_rows_text: str = "",
    table_rows_file: Path | None = None,
    source_path: Path | None = None,
    background_path: Path | None = None,
    logo_path: Path | None = None,
    output_dir: Path | None = None,
) -> tuple[list[Path], ManualReportInput]:
    output_dir = output_dir or get_output_dir()
    output_dir.mkdir(parents=True, exist_ok=True)

    tier_metrics = [MetricPoint(label=label, value=int(value), raw_value=str(int(value))) for label, value in tier_values.items()]
    media_metrics = [MetricPoint(label=label, value=int(value), raw_value=str(int(value))) for label, value in media_values.items()]
    month_labels = ["Mes 1", "Mes 2", "Mes 3", "Mes 4", "Mes 5", "Mes 6", "Mes 7", "Mes 8", "Mes 9", "Mes 10", "Mes 11", "Mes 12"]
    monthly_trend = [
        MetricPoint(label=label, value=int(value), raw_value=str(int(value)))
        for label, value in zip(month_labels, monthly_values)
    ]
    manual_input = ManualReportInput(
        title=report_title,
        client_name=client_name,
        report_month=report_month,
        monthly_summary=monthly_summary,
        executive_comment=executive_comment,
        next_steps=next_steps,
        tier_metrics=tier_metrics,
        media_metrics=media_metrics,
        monthly_trend=monthly_trend,
        reach_value=reach_value,
        valuation_value=valuation_value,
        table_rows=parse_table_file(table_rows_file) if table_rows_file else parse_table_rows(table_rows_text),
    )

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    cover_data = ReportData(
        source_text="",
        title=report_title,
        source_name=source_path.name if source_path else "manual",
        client_name=client_name,
        report_month=report_month,
    )
    add_cover_slide(prs, cover_data, background_path, logo_path)
    add_month_summary_slide(prs, monthly_summary, background_path, logo_path)

    add_named_chart_slide(
        prs,
        ChartRequest(title="Distribución de Tiers", chart_type="bar", aliases=(), metrics=tier_metrics),
        background_path,
        logo_path,
    )
    add_named_chart_slide(
        prs,
        ChartRequest(title="Distribución de Medios", chart_type="bar", aliases=(), metrics=media_metrics),
        background_path,
        logo_path,
    )
    add_monthly_trend_slide(prs, monthly_trend, background_path, logo_path)
    add_scope_slide(prs, reach_value, valuation_value, background_path, logo_path)

    exec_data = ReportData(
        source_text="",
        title=report_title,
        source_name=source_path.name if source_path else "manual",
        client_name=client_name,
        report_month=report_month,
        executive_comment=executive_comment,
        next_steps=next_steps,
    )
    add_next_steps_slide(prs, exec_data, background_path, logo_path)
    add_results_table_slide(prs, manual_input.table_rows, background_path, logo_path)

    source_stem = source_path.stem if source_path else client_name or "reporte"
    slug = re.sub(r"[^a-z0-9_]+", "_", source_stem.lower()).strip("_") or "reporte"
    output_path = output_dir / f"Reporte_Automatico_{slug}_{datetime.now().year}.pptx"
    prs.save(output_path)
    write_manual_report_summary(output_path, manual_input, background_path, logo_path)
    return [output_path], manual_input


def generate_report_from_pdf(
    pdf_path: Path,
    executive_comment: str = "",
    next_steps: str = "",
    report_title: str | None = None,
    client_name: str = "",
    report_month: str = "",
    background_path: Path | None = None,
    logo_path: Path | None = None,
    output_dir: Path | None = None,
) -> tuple[list[Path], ReportData]:
    output_dir = output_dir or get_output_dir()
    output_dir.mkdir(parents=True, exist_ok=True)

    data = parse_report_pdf(
        pdf_path=pdf_path,
        executive_comment=executive_comment,
        next_steps=next_steps,
        report_title=report_title,
        client_name=client_name,
        report_month=report_month,
    )

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_cover_slide(prs, data, background_path, logo_path)

    for chart in data.requested_charts:
        add_named_chart_slide(prs, chart, background_path, logo_path)

    add_monthly_trend_slide(prs, data.monthly_trend, background_path, logo_path)
    add_exec_slide(prs, data, background_path, logo_path)
    add_next_steps_slide(prs, data, background_path, logo_path)

    slug = re.sub(r"[^a-z0-9_]+", "_", pdf_path.stem.lower()).strip("_") or "reporte"
    output_path = output_dir / f"Reporte_Automatico_{slug}_{datetime.now().year}.pptx"
    prs.save(output_path)
    write_report_summary(output_path, data, background_path, logo_path)
    return [output_path], data
